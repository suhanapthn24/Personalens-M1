import json

import pytest

from personalens.ingestion.parsers import parse_file


def test_docx_headings_tables_in_order(tmp_path):
    from docx import Document

    d = Document()
    d.add_heading("Bayes Theorem", level=1)
    d.add_paragraph("Posterior equals likelihood times prior over evidence.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Term", "Meaning"
    t.cell(1, 0).text, t.cell(1, 1).text = "Prior", "Belief before data"
    d.add_heading("Naive Bayes", level=1)
    d.add_paragraph("Assumes feature independence.")
    p = tmp_path / "n.docx"
    d.save(p)

    segs = parse_file(p, ".docx").segments
    assert [s.section for s in segs] == ["Bayes Theorem", "Naive Bayes"]
    assert "Prior | Belief before data" in segs[0].text
    assert segs[0].text.index("Posterior") < segs[0].text.index("Prior |")


def test_pptx_titles_bullets_notes(tmp_path):
    from pptx import Presentation

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Conditional Probability"
    s.placeholders[1].text = "P(A|B) = P(A and B) / P(B)\nDefined when P(B) > 0"
    s.notes_slide.notes_text_frame.text = "Stress the P(B) > 0 condition"
    prs.slides.add_slide(prs.slide_layouts[6])  # blank slide -> skipped
    p = tmp_path / "s.pptx"
    prs.save(p)

    segs = parse_file(p, ".pptx").segments
    assert len(segs) == 1
    assert segs[0].page == 1 and segs[0].section == "Conditional Probability"
    assert "Speaker notes: Stress" in segs[0].text


def test_pdf_pages_and_scanned_warning(tmp_path):
    import pymupdf

    doc = pymupdf.open()
    for i in range(1, 3):
        page = doc.new_page()
        page.insert_text((72, 100), f"Page {i} discusses gradient descent and learning rates in depth. " * 2, fontsize=8)
    p = tmp_path / "a.pdf"
    doc.save(p)
    parsed = parse_file(p, ".pdf")
    assert [s.page for s in parsed.segments] == [1, 2]
    assert "gradient descent" in parsed.segments[0].text
    assert parsed.warnings == []

    blank = pymupdf.open()
    blank.new_page()
    q = tmp_path / "blank.pdf"
    blank.save(q)
    assert parse_file(q, ".pdf").warnings  # scanned/empty warning


def test_markdown_sections_ignore_fenced_hashes(tmp_path):
    p = tmp_path / "n.md"
    p.write_text("# Intro\nhello world\n```\n# not a heading\n```\n## Next\nmore text\n")
    segs = parse_file(p, ".md").segments
    assert [s.section for s in segs] == ["Intro", "Next"]


def test_notebook_cells(tmp_path):
    nb = {"cells": [
        {"cell_type": "markdown", "source": ["# Title\n", "some prose"]},
        {"cell_type": "code", "source": "x = 1", "outputs": [{"data": "huge"}]},
        {"cell_type": "raw", "source": "skip"},
        {"cell_type": "code", "source": "   "},
    ]}
    p = tmp_path / "n.ipynb"
    p.write_text(json.dumps(nb))
    segs = parse_file(p, ".ipynb").segments
    assert [s.kind for s in segs] == ["prose", "code"]


def test_non_utf8_text_falls_back(tmp_path):
    p = tmp_path / "n.txt"
    p.write_bytes("caf\xe9 society".encode("cp1252"))
    assert "caf\u00e9" in parse_file(p, ".txt").segments[0].text


def test_unsupported_extension():
    with pytest.raises(ValueError):
        parse_file("x.exe", ".exe")
