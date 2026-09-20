"""Stage 1a: turn an uploaded file into text Segments (with page / section metadata)."""
from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from pathlib import Path

CODE_EXTS = {".py", ".ipynb", ".js", ".ts", ".java", ".c", ".cpp", ".h", ".cs", ".go", ".rs", ".r", ".sql", ".sh"}
NOTE_EXTS = {".docx", ".txt", ".md", ".markdown"}
MARKDOWN_EXTS = {".md", ".markdown"}


@dataclass
class Segment:
    text: str
    page: int | None = None  # PDF page / slide number
    section: str | None = None  # heading or slide title
    kind: str = "prose"  # "prose" | "code"
    unwrap: bool = False  # True when hard-wrapped lines need re-joining (PDF)


@dataclass
class ParsedDocument:
    segments: list[Segment]
    warnings: list[str] = field(default_factory=list)


# ---------------------------------------------------------------- PDF
def parse_pdf(path: Path) -> ParsedDocument:
    try:
        import pymupdf as fitz
    except ImportError:  # older PyMuPDF
        import fitz

    segments: list[Segment] = []
    with fitz.open(str(path)) as doc:
        if doc.needs_pass:
            raise ValueError("PDF is password-protected")
        for i, page in enumerate(doc, start=1):
            segments.append(Segment(page.get_text("text", sort=True), page=i, unwrap=True))

    warnings = []
    if segments and sum(len(s.text.strip()) for s in segments) < 50 * len(segments):
        warnings.append(
            "Very little extractable text: this PDF may be scanned images. OCR is not supported yet."
        )
    return ParsedDocument(segments, warnings)


# ---------------------------------------------------------------- DOCX
def _table_rows(table) -> list[str]:
    rows = []
    for row in table.rows:
        cells: list[str] = []
        for cell in row.cells:
            t = cell.text.strip().replace("\n", " ")
            if t and (not cells or cells[-1] != t):  # merged cells repeat their text
                cells.append(t)
        if cells:
            rows.append(" | ".join(cells))
    return rows


def parse_docx(path: Path) -> ParsedDocument:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(str(path))
    segments: list[Segment] = []
    buf: list[str] = []
    section: str | None = None

    def flush() -> None:
        if buf:
            segments.append(Segment("\n".join(buf), section=section))
            buf.clear()

    for child in doc.element.body.iterchildren():  # body order keeps tables in place
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if style.startswith("Heading") or style == "Title":
                flush()
                section = text
            buf.append(text)
        elif tag == "tbl":
            buf.extend(_table_rows(Table(child, doc)))
    flush()
    return ParsedDocument(segments)


# ---------------------------------------------------------------- PPTX
def _collect_shape_text(shapes, lines: list[str]) -> None:
    from pptx.shapes.group import GroupShape

    for sh in shapes:
        if isinstance(sh, GroupShape):
            _collect_shape_text(sh.shapes, lines)
            continue
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                if para.text.strip():
                    lines.append(para.text.strip())
        if getattr(sh, "has_table", False) and sh.has_table:
            lines.extend(_table_rows(sh.table))


def parse_pptx(path: Path) -> ParsedDocument:
    from pptx import Presentation

    prs = Presentation(str(path))
    segments: list[Segment] = []
    for i, slide in enumerate(prs.slides, start=1):
        lines: list[str] = []
        _collect_shape_text(slide.shapes, lines)

        title = None
        t_shape = slide.shapes.title
        if t_shape is not None and t_shape.has_text_frame:
            title = t_shape.text_frame.text.strip() or None

        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"Speaker notes: {notes}")

        if lines:
            segments.append(Segment("\n".join(lines), page=i, section=title))
    return ParsedDocument(segments)


# ---------------------------------------------------------------- text / markdown / code
def _read_text(path: Path) -> str:
    raw = Path(path).read_bytes()
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace")


def parse_text(path: Path, ext: str) -> ParsedDocument:
    text = _read_text(path)
    if ext not in MARKDOWN_EXTS:
        return ParsedDocument([Segment(text)])

    segments: list[Segment] = []
    buf: list[str] = []
    section: str | None = None
    in_fence = False

    def flush() -> None:
        if any(line.strip() for line in buf):
            segments.append(Segment("\n".join(buf), section=section))
        buf.clear()

    for line in text.splitlines():
        if line.lstrip().startswith("```"):
            in_fence = not in_fence
        m = None if in_fence else re.match(r"^#{1,6}\s+(.*\S)\s*$", line)
        if m:
            flush()
            section = m.group(1)
        buf.append(line)
    flush()
    return ParsedDocument(segments)


def parse_notebook(path: Path) -> ParsedDocument:
    nb = json.loads(_read_text(path))
    segments: list[Segment] = []
    for cell in nb.get("cells", []):
        ctype = cell.get("cell_type")
        if ctype not in {"code", "markdown"}:
            continue
        src = cell.get("source", "")
        src = "".join(src) if isinstance(src, list) else src
        if src.strip():  # outputs are dropped on purpose (can be huge / binary)
            segments.append(Segment(src, kind="code" if ctype == "code" else "prose"))
    return ParsedDocument(segments)


def parse_code(path: Path) -> ParsedDocument:
    return ParsedDocument([Segment(_read_text(path), kind="code")])


# ---------------------------------------------------------------- dispatcher
def parse_file(path: Path, ext: str) -> ParsedDocument:
    ext = ext.lower()
    if ext == ".pdf":
        return parse_pdf(path)
    if ext == ".pptx":
        return parse_pptx(path)
    if ext == ".docx":
        return parse_docx(path)
    if ext in NOTE_EXTS:
        return parse_text(path, ext)
    if ext == ".ipynb":
        return parse_notebook(path)
    if ext in CODE_EXTS:
        return parse_code(path)
    raise ValueError(f"Unsupported file type: {ext}")
